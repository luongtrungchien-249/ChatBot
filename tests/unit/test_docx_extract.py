from pathlib import Path

import docx

from knowledge.ingest.extract import extract_document


def test_docx_keeps_tables_in_document_order_and_preserves_heading(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_heading("Synthetic policy", 1)
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Allowance"
    table.cell(1, 0).text = "Employee"
    table.cell(1, 1).text = "12 days"
    document.add_paragraph("After table")
    path = tmp_path / "synthetic.docx"
    document.save(str(path))
    extracted = extract_document(path)
    assert extracted.text == (
        "# Synthetic policy\n\nBefore table\n\nRole | Allowance\nEmployee | 12 days"
        "\n\nAfter table"
    )
    assert all(page is None for _, page in extracted.ban_do_trang)


def test_docx_nested_table_and_merged_cells_are_not_silently_lost(tmp_path: Path) -> None:
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
    table.cell(1, 0).text = "Limit"
    nested = table.cell(1, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Travel"
    nested.cell(0, 1).text = "42"
    path = tmp_path / "nested.docx"
    document.save(str(path))
    text = extract_document(path).text
    assert text.count("Merged heading") == 1
    assert "Limit | Travel | 42" in text
